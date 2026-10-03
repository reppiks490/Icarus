"""Parity check: this engine's trades vs a TradingView Strategy Tester "List of Trades" export.

How to produce the TradingView side:
  1. Load THE_PULSE_OF_ICARUS v3/v3.1 on the SAME symbol and timeframe the engine runs
     (e.g. COINBASE:BTCUSD, 5m) with the same input changes (crypto profile:
     `icarus-engine inputs --profile crypto`).
  2. Strategy Tester → List of Trades → Export (CSV).
  3. `icarus-engine parity --asset BTC --tf 5 --warmup 2000 --tv-csv "<file>"`

Matching: trades are grouped by entry (TradingView lists each partial exit as a
row pair with the same trade number). An engine trade matches a TV trade when
direction agrees and the entry time is within `tol` chart bars, after the
chart-timezone offset that maximises matches is chosen automatically.
"""
from __future__ import annotations

import csv
import math
import re
from collections import defaultdict
from datetime import datetime, timezone
from typing import Any, Dict, List

from .runtime import AssetRunner


def _text_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        dt = value
        if dt.tzinfo is not None:
            dt = dt.astimezone(timezone.utc).replace(tzinfo=None)
        return dt.strftime("%Y-%m-%d %H:%M:%S")
    return str(value).strip()


def _num(value: Any) -> float:
    if isinstance(value, bool):
        return float("nan")
    if isinstance(value, (int, float)):
        return float(value)
    s = _text_value(value).replace(",", "").replace("$", "").replace("%", "").strip()
    try:
        return float(s)
    except ValueError:
        return float("nan")


def _parse_dt(value: Any) -> int:
    if isinstance(value, datetime):
        dt = value
        if dt.tzinfo is None or dt.utcoffset() is None:
            dt = dt.replace(tzinfo=timezone.utc)
        else:
            dt = dt.astimezone(timezone.utc)
        return int(dt.timestamp())
    s = _text_value(value)
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M:%SZ", "%m/%d/%Y %H:%M", "%d.%m.%Y %H:%M"):
        try:
            return int(datetime.strptime(s, fmt).replace(tzinfo=timezone.utc).timestamp())
        except ValueError:
            pass
    raise ValueError(f"unrecognised date: {s!r}")


def _trade_columns(keys: List[str]) -> Dict[str, str]:
    lookup = {str(k).strip().lower(): str(k) for k in keys if str(k or "").strip()}

    def col(*names: str) -> str:
        for name in names:
            for lowered, original in lookup.items():
                if name in lowered:
                    return original
        raise KeyError(names)

    return {
        "no": col("trade #", "trade number", "trade"),
        "type": col("type"),
        "dt": col("date"),
        "signal": col("signal"),
        "price": col("price"),
        "qty": col("position size", "size (qty)", "contracts", "qty"),
        "pnl": col("net p&l usd", "net pnl usd", "net pnl", "profit usd", "profit", "p&l"),
    }


def _validate_historical_trade_rows(rows: List[Dict[str, Any]], columns: Dict[str, str]) -> None:
    """Validate row relationships without assuming export order is chronological."""
    positions: Dict[str, Dict[str, Any]] = {}
    for index, row in enumerate(rows, start=1):
        no = _text_value(row.get(columns['no']))
        typ = _text_value(row.get(columns['type'])).lower()
        if None in row or not no or typ not in {'entry long', 'entry short', 'exit long', 'exit short'}:
            raise ValueError(f'trade row {index} has missing identity or unsupported type')
        qty = _num(row.get(columns['qty']))
        if not math.isfinite(qty) or qty <= 0:
            raise ValueError(f'trade row {index} quantity must be finite and positive')
        role, direction = typ.split()
        position = positions.setdefault(no, {'entry': None, 'exits': []})
        is_open = role == 'exit' and (
            _text_value(row.get(columns['dt'])).lower() == 'open'
            or _text_value(row.get(columns['signal'])).lower() == 'open'
        )
        timestamp = None if is_open else _parse_dt(row.get(columns['dt']))
        if not is_open and not math.isfinite(_num(row.get(columns['price']))):
            raise ValueError(f'trade row {index} price must be finite')
        if role == 'exit' and not is_open and not math.isfinite(_num(row.get(columns['pnl']))):
            raise ValueError(f'trade row {index} realized PnL must be finite')
        if role == 'entry':
            if position['entry'] is not None:
                raise ValueError(f'trade {no} has repeated entry rows')
            position['entry'] = (direction, timestamp)
        else:
            position['exits'].append((direction, timestamp))
    for no, position in positions.items():
        if position['entry'] is None:
            raise ValueError(f'trade {no} has an exit without an entry')
        direction, entry_timestamp = position['entry']
        for exit_direction, exit_timestamp in position['exits']:
            if exit_direction != direction:
                raise ValueError(f'trade {no} entry and exit directions disagree')
            if exit_timestamp is not None and exit_timestamp < entry_timestamp:
                raise ValueError(f'trade {no} exit precedes its entry')


def _read_tv_trade_rows(rows: List[Dict[str, Any]], *, strict: bool = False) -> List[Dict[str, Any]]:
    if not rows:
        return []
    columns = _trade_columns(list(rows[0].keys()))
    if strict:
        _validate_historical_trade_rows(rows, columns)
    by_no: Dict[str, Dict[str, Any]] = {}
    for row in rows:
        no = _text_value(row.get(columns["no"]))
        typ = _text_value(row.get(columns["type"])).lower()
        if not no or not typ:
            continue
        trade = by_no.setdefault(no, {"no": no, "dir": 0, "pieces": []})
        if typ.startswith("entry"):
            trade["dir"] = 1 if "long" in typ else -1
            trade["entry_ts"] = _parse_dt(row.get(columns["dt"]))
            trade["entry_px"] = _num(row.get(columns["price"]))
            trade["entry_sig"] = _text_value(row.get(columns["signal"]))
            trade["qty"] = _num(row.get(columns["qty"]))
        elif typ.startswith("exit"):
            dt_text = _text_value(row.get(columns["dt"])).lower()
            sig = _text_value(row.get(columns["signal"]))
            if dt_text == "open" or sig.lower() == "open":
                trade["pieces"].append({
                    "ts": None,
                    "px": float("nan"),
                    "sig": "OPEN",
                    "qty": _num(row.get(columns["qty"])),
                    "pnl": _num(row.get(columns["pnl"])),
                })
                continue
            trade["pieces"].append({
                "ts": _parse_dt(row.get(columns["dt"])),
                "px": _num(row.get(columns["price"])),
                "sig": sig,
                "qty": _num(row.get(columns["qty"])),
                "pnl": _num(row.get(columns["pnl"])),
            })

    grouped: Dict[tuple, Dict[str, Any]] = {}
    for trade in by_no.values():
        if trade.get("entry_ts") is None or (not strict and not trade.get("entry_ts")):
            continue
        key = (trade["entry_ts"], trade["dir"], round(trade["entry_px"], 4))
        group = grouped.setdefault(key, {
            "no": trade["no"],
            "dir": trade["dir"],
            "entry_ts": trade["entry_ts"],
            "entry_px": trade["entry_px"],
            "entry_sig": trade["entry_sig"],
            "qty": 0.0,
            "pieces": [],
        })
        group["qty"] += trade["qty"]
        group["pieces"].extend(trade["pieces"])
    out = list(grouped.values())
    out.sort(key=lambda trade: trade["entry_ts"])
    return out


def read_tv_trades(path: str, *, strict: bool = False) -> List[Dict[str, Any]]:
    with open(path, "r", encoding="utf-8-sig", newline="") as fh:
        return read_tv_trades_text(fh.read(), strict=strict)


def read_tv_trades_text(text: str, *, strict: bool = False) -> List[Dict[str, Any]]:
    import io as _io
    reader = csv.DictReader(_io.StringIO(text.lstrip("﻿")))
    headers = [str(value).strip().lower() for value in reader.fieldnames or []]
    if strict and len(headers) != len(set(headers)):
        raise ValueError('historical trade table has duplicate column names')
    return _read_tv_trade_rows(list(reader), strict=strict)


def read_tv_trades_xlsx(path: str, *, strict: bool = False) -> List[Dict[str, Any]]:
    """Read a TradingView Strategy Tester trade table from an XLSX workbook.

    The sheet name and header row are discovered from required TradingView
    columns. Exactly one trade table must be present; ambiguity fails closed.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    candidates: List[List[Dict[str, Any]]] = []
    try:
        for worksheet in workbook.worksheets:
            matrix = [list(row) for row in worksheet.iter_rows(values_only=True)]
            for index, values in enumerate(matrix[:100]):
                headers = [_text_value(value) for value in values]
                if not any(headers):
                    continue
                try:
                    _trade_columns(headers)
                except KeyError:
                    continue
                named_headers = [header.lower() for header in headers if header]
                if strict and len(named_headers) != len(set(named_headers)):
                    raise ValueError('historical trade table has duplicate column names')
                rows: List[Dict[str, Any]] = []
                for raw in matrix[index + 1:]:
                    if not any(value not in (None, "") for value in raw):
                        continue
                    record = {
                        header: (raw[pos] if pos < len(raw) else None)
                        for pos, header in enumerate(headers)
                        if header
                    }
                    rows.append(record)
                candidates.append(rows)
                break
    finally:
        workbook.close()

    if not candidates:
        raise ValueError("no TradingView trade table found in XLSX workbook")
    if len(candidates) != 1:
        raise ValueError("multiple TradingView trade tables found in XLSX workbook")
    return _read_tv_trade_rows(candidates[0], strict=strict)


def engine_trades(r: AssetRunner) -> List[Dict[str, Any]]:
    groups: Dict[tuple, Dict[str, Any]] = {}
    for ct in r.em.closed:
        k = (ct.entry_id, ct.entry_ts)
        g = groups.setdefault(k, {"id": ct.entry_id, "dir": ct.direction, "entry_ts": ct.entry_ts, "entry_px": ct.entry_price, "qty": 0, "pieces": []})
        g["qty"] += ct.qty
        g["pieces"].append({"ts": ct.exit_ts, "px": ct.exit_price, "sig": ct.exit_comment, "qty": ct.qty, "pnl": ct.profit})
    for ot in r.em.open:                                       # still-open entries take part in the matching too
        k = (ot.entry_id, ot.entry_ts)
        g = groups.setdefault(k, {"id": ot.entry_id, "dir": ot.direction, "entry_ts": ot.entry_ts, "entry_px": ot.entry_price, "qty": 0, "pieces": []})
        g["qty"] += ot.qty
        g["pieces"].append({"ts": None, "px": float("nan"), "sig": "OPEN", "qty": ot.qty, "pnl": 0.0})
    out = list(groups.values())
    out.sort(key=lambda t: t["entry_ts"])
    return out


def engine_trades_from_rows(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """The same structure from a backtest job's trade rows (backtest.py)."""
    groups: Dict[tuple, Dict[str, Any]] = {}
    for t in rows:
        d = 1 if t["type"] == "long" else -1
        k = (t.get("entry_signal") or t["type"], t["entry_ts"])
        g = groups.setdefault(k, {"id": t.get("entry_signal") or t["type"], "dir": d, "entry_ts": t["entry_ts"], "entry_px": t["entry_px"], "qty": 0, "pieces": []})
        g["qty"] += t["qty"]
        g["pieces"].append({"ts": t["exit_ts"], "px": (t["exit_px"] if t["exit_px"] is not None else float("nan")), "sig": ("OPEN" if t.get("open") else t["exit_signal"]), "qty": t["qty"], "pnl": t["pnl"] if not t.get("open") else 0.0})
    out = list(groups.values())
    out.sort(key=lambda t: t["entry_ts"])
    return out


def compare(r: AssetRunner, tv_csv: str, tol_bars: int = 1) -> Dict[str, Any]:
    return compare_lists(engine_trades(r), read_tv_trades(tv_csv), r.chart_minutes * 60, tol_bars)


def compare_lists(eng: List[Dict[str, Any]], tv: List[Dict[str, Any]], bar: int, tol_bars: int = 1) -> Dict[str, Any]:
    if not tv or not eng:
        return {"summary": {"tv_trades": len(tv), "engine_trades": len(eng), "note": "nothing to compare"}, "lines": []}
    # restrict to the overlapping window
    lo = max(min(t["entry_ts"] for t in eng), min(t["entry_ts"] for t in tv) - 14 * 3600)
    best = None
    for off_half in range(-28, 29):
        off = off_half * 1800
        used = set(); n = 0
        for e in eng:
            for j, t in enumerate(tv):
                if j in used or t["dir"] != e["dir"]:
                    continue
                if abs((t["entry_ts"] + off) - e["entry_ts"]) <= tol_bars * bar:
                    used.add(j); n += 1; break
        if best is None or n > best[1]:
            best = (off, n)
    off, _ = best
    tv_adj = [dict(t, entry_ts=t["entry_ts"] + off, pieces=[dict(p, ts=(p["ts"] + off) if p["ts"] is not None else None) for p in t["pieces"]]) for t in tv]
    tv_w = [t for t in tv_adj if t["entry_ts"] >= lo]
    eng_w = [t for t in eng if t["entry_ts"] >= lo]
    used = set(); matched = []; unmatched_eng = []
    for e in eng_w:
        hit = None
        for j, t in enumerate(tv_w):
            if j in used or t["dir"] != e["dir"]:
                continue
            if abs(t["entry_ts"] - e["entry_ts"]) <= tol_bars * bar:
                hit = j; break
        if hit is None:
            unmatched_eng.append(e)
        else:
            used.add(hit); matched.append((e, tv_w[hit]))
    unmatched_tv = [t for j, t in enumerate(tv_w) if j not in used]
    px_diff = [abs(e["entry_px"] - t["entry_px"]) / t["entry_px"] * 1e4 for e, t in matched if t["entry_px"] == t["entry_px"]]
    exit_agree = sum(1 for e, t in matched if sorted(p["sig"] for p in e["pieces"]) == sorted(p["sig"] for p in t["pieces"]))
    pnl_eng = sum(p["pnl"] for e in eng_w for p in e["pieces"]); pnl_tv = sum(p["pnl"] for t in tv_w for p in t["pieces"] if p["pnl"] == p["pnl"])
    lines = []
    for e, t in matched:
        lines.append(f"MATCH  {_ts(e['entry_ts'])} {'L' if e['dir']>0 else 'S'} eng {e['entry_px']:.6g} tv {t['entry_px']:.6g}  exits eng {[p['sig'] for p in e['pieces']]} tv {[p['sig'] for p in t['pieces']]}")
    for e in unmatched_eng:
        lines.append(f"ENGINE-ONLY {_ts(e['entry_ts'])} {'L' if e['dir']>0 else 'S'} @ {e['entry_px']:.6g} exits {[p['sig'] for p in e['pieces']]}")
    for t in unmatched_tv:
        lines.append(f"TV-ONLY     {_ts(t['entry_ts'])} {'L' if t['dir']>0 else 'S'} @ {t['entry_px']:.6g} exits {[p['sig'] for p in t['pieces']]}")
    n_all = len(eng_w) + len(unmatched_tv)
    return {"summary": {
        "window_start": _ts(lo), "tz_offset_hours": off / 3600, "engine_trades": len(eng_w), "tv_trades": len(tv_w),
        "matched": len(matched), "engine_only": len(unmatched_eng), "tv_only": len(unmatched_tv),
        "match_rate_pct": round(len(matched) / n_all * 100, 1) if n_all else None,
        "exit_signature_agreement_pct": round(exit_agree / len(matched) * 100, 1) if matched else None,
        "avg_entry_price_diff_bps": round(sum(px_diff) / len(px_diff), 2) if px_diff else None,
        "net_pnl_engine": round(pnl_eng, 2), "net_pnl_tv": round(pnl_tv, 2),
    }, "lines": lines}


def _ts(t: int) -> str:
    return datetime.fromtimestamp(int(t), tz=timezone.utc).strftime("%Y-%m-%d %H:%M")
